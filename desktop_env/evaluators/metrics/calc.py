
import logging
import openpyxl
from typing import Any, Dict, Union

logger = logging.getLogger("desktopenv.metrics.calc")

def check_excel_value(file_path: str, rule: Dict[str, Any]) -> float:
    """
    Check if a specific cell in an Excel file matches an expected value.
    
    Args:
        file_path (str): Path to the Excel file.
        rule (dict): Rule containing:
            - sheet (str): Sheet name (default: active sheet).
            - cell (str): Cell coordinate (e.g., "A1").
            - value (Any): Expected value.
            - type (str, optional): "currency" to check checks if value is float/int.
    
    Returns:
        float: 1.0 if match, 0.0 otherwise.
    """
    if not file_path:
        logger.error("check_excel_value: file_path is None")
        return 0.0
        
    try:
        # Load with data_only=True to get the calculate results
        wb = openpyxl.load_workbook(file_path, data_only=True)
        
        sheet_name = rule.get("sheet")
        if sheet_name:
            if sheet_name not in wb.sheetnames:
                logger.error(f"Sheet '{sheet_name}' not found in {file_path}")
                return 0.0
            ws = wb[sheet_name]
        else:
            ws = wb.active
            
        cell_coord = rule.get("cell")
        if not cell_coord:
            logger.error("check_excel_value: 'cell' not specified in rule")
            return 0.0
            
        actual_value = ws[cell_coord].value
        expected_value = rule.get("value")
        check_type = rule.get("type")
        
        logger.info(f"Checking cell {cell_coord}: Actual={actual_value}, Expected={expected_value} (Type={check_type})")
        
        if check_type == "currency":
            # Just check if it's a number
            if isinstance(actual_value, (int, float)):
                return 1.0
            else:
                return 0.0
                
        # Numeric Comparison
        try:
            float_actual = float(actual_value)
            float_expected = float(expected_value)
            if abs(float_actual - float_expected) < 0.01:
                return 1.0
        except (ValueError, TypeError):
            # Fallback to string comparison
            pass
            
        # String/Equality check
        if str(actual_value).strip() == str(expected_value).strip():
            return 1.0
            
        return 0.0
        
    except Exception as e:
        logger.error(f"Error in check_excel_value: {e}")
        return 0.0

def check_excel_formula(file_path: str, rule: Dict[str, Any]) -> float:
    """
    Check if a specific cell contains a formula starting with a string.
    
    Args:
        file_path (str): Path to the Excel file.
        rule (dict): Rule containing:
            - sheet (str): Sheet name.
            - cell (str): Cell coordinate.
            - formula_starts_with (str): Expected prefix (e.g., "=SUM").
    """
    if not file_path:
        return 0.0
        
    try:
        # Load with data_only=False to get formulas
        wb = openpyxl.load_workbook(file_path, data_only=False)
        
        sheet_name = rule.get("sheet")
        if sheet_name and sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            ws = wb.active
            
        cell_coord = rule.get("cell")
        if not cell_coord:
            return 0.0
            
        cell_content = ws[cell_coord].value
        expected_prefix = rule.get("formula_starts_with", "=")
        
        logger.info(f"Checking formula {cell_coord}: Content={cell_content}, Expected Prefix={expected_prefix}")
        
        if isinstance(cell_content, str) and cell_content.upper().startswith(expected_prefix.upper()):
            return 1.0
            
        return 0.0
        
    except Exception as e:
        logger.error(f"Error in check_excel_formula: {e}")
        return 0.0
